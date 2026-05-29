# ============================================================
# utils/excel_generator.py — Excel Report Generator
# Uses openpyxl to produce styled .xlsx files
# ============================================================
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Style constants ──────────────────────────────────────────
HEADER_FILL   = PatternFill('solid', fgColor='1E3A5F')
HEADER_FONT   = Font(bold=True, color='FFFFFF', size=11)
ALT_FILL      = PatternFill('solid', fgColor='EBF0F7')
PRESENT_FILL  = PatternFill('solid', fgColor='D4EDDA')
LATE_FILL     = PatternFill('solid', fgColor='FFF3CD')
ABSENT_FILL   = PatternFill('solid', fgColor='F8D7DA')
THIN_BORDER   = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)


def _apply_header(ws, headers, row=1):
    """Write a styled header row."""
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border    = THIN_BORDER


def _auto_width(ws):
    """Auto-fit column widths based on content."""
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)


def generate_attendance_excel(subject, records) -> bytes:
    """
    Generate an attendance report Excel file.

    Args:
        subject:  Subject ORM object
        records:  list of Attendance ORM objects

    Returns:
        bytes — the .xlsx file content
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Attendance Report'

    # ── Title block ──────────────────────────────────────────
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value     = f'Attendance Report — {subject.name} ({subject.code})'
    title_cell.font      = Font(bold=True, size=14, color='1E3A5F')
    title_cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:G2')
    ws['A2'].value     = f'Semester {subject.semester}  |  Generated on: {__import__("datetime").date.today()}'
    ws['A2'].alignment = Alignment(horizontal='center')
    ws['A2'].font      = Font(italic=True, size=10, color='666666')

    # ── Header ───────────────────────────────────────────────
    headers = ['#', 'Reg Number', 'Student Name', 'Date', 'Time In', 'Status', 'Confidence']
    _apply_header(ws, headers, row=4)

    # ── Data rows ────────────────────────────────────────────
    for idx, rec in enumerate(records, 1):
        row_num = idx + 4
        student = rec.student   # backref

        values = [
            idx,
            student.reg_number,
            student.name,
            str(rec.date),
            str(rec.time_in.strftime('%H:%M') if rec.time_in else '—'),
            rec.status.value,
            f'{rec.confidence_score:.2f}' if rec.confidence_score else '—',
        ]

        # Row background colour by status
        status_fills = {
            'Present': PRESENT_FILL,
            'Late':    LATE_FILL,
            'Absent':  ABSENT_FILL,
        }
        row_fill = status_fills.get(rec.status.value, ALT_FILL if idx % 2 == 0 else None)

        for col, val in enumerate(values, 1):
            cell           = ws.cell(row=row_num, column=col, value=val)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal='center')
            if row_fill:
                cell.fill = row_fill

    # ── Summary row ──────────────────────────────────────────
    summary_row = len(records) + 6
    total   = len(records)
    present = sum(1 for r in records if r.status.value in ('Present', 'Late'))
    ws.cell(row=summary_row, column=1, value='Total').font = Font(bold=True)
    ws.cell(row=summary_row, column=2, value=total)
    ws.cell(row=summary_row, column=3, value=f'Present: {present} ({round(present/total*100,1) if total else 0}%)')

    _auto_width(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_student_report_excel(students, dept_name: str) -> bytes:
    """Generate a full student list report."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Student Report'

    ws.merge_cells('A1:H1')
    ws['A1'].value     = f'Student Report — {dept_name}'
    ws['A1'].font      = Font(bold=True, size=14, color='1E3A5F')
    ws['A1'].alignment = Alignment(horizontal='center')

    headers = ['#', 'Reg Number', 'Name', 'Email', 'Semester', 'Year', 'Attendance %', 'Risk Level']
    _apply_header(ws, headers, row=3)

    for idx, s in enumerate(students, 1):
        row = idx + 3
        attend_pct = s.get_attendance_percentage()
        values     = [idx, s.reg_number, s.name, s.email, s.semester, s.year, f'{attend_pct}%', '—']
        for col, val in enumerate(values, 1):
            cell           = ws.cell(row=row, column=col, value=val)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal='center')
            if idx % 2 == 0:
                cell.fill = ALT_FILL

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
